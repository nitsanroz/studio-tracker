#!/usr/bin/env python3
"""
One-time import of the "Estimations / Pricing" workbook into Leads estimates
(0045). One estimate per tab; tabs that are later versions of the same quote
become v2/v3 on the same lead.

    python3 scripts/import-estimates.py <workbook.xlsx>            # dry run
    python3 scripts/import-estimates.py <workbook.xlsx> --write    # writes

Needs NEXT_PUBLIC_SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY in the env
(`set -a; . ./.env.local; set +a` before running).

Decisions (Nitsan, 2026-10-03):
  * every tab keeps ITS OWN hourly rate (300 for the older quotes, 350 newer);
  * Mobile / Prep for dev / QA rows become live percentage lines where the tab
    states the percentage (0.1, 0.15, "(10%)"), fixed hours otherwise;
  * tabs with no matching lead get a NEW lead, its contact traced from Gmail
    headers (passed in below as TRACED — nothing here reads mail);
  * "Monday" is a days x hours calculation, not an estimate, and is skipped.

Imported estimates are APPROVED (read-only history). Copying one from the
"New estimate" picker gives an editable draft. Re-runnable: an estimate is
skipped when its lead already has one carrying the same import note.

⚠️ THE CHECK THAT MATTERS: every tab's computed total is compared with the
tab's own "Total in work Hours" row, and any gap is printed. A parsing slip
shows up as a mismatch here, not as a wrong quote six months later.
"""
import json
import os
import re
import sys
import urllib.request
from datetime import datetime, timezone

import openpyxl

WRITE = "--write" in sys.argv
# --only=<tab>: restrict to one tab (restoring a single deleted lead's estimate).
ONLY = next((a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith("--only=")), None)
PATH = next(a for a in sys.argv[1:] if not a.startswith("--"))
URL = os.environ.get("NEXT_PUBLIC_SUPABASE_URL", "").strip('"')
KEY = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "").strip('"')

SKIP_TABS = {"Monday"}

# tab -> existing lead company (exact, case-insensitive) when the name differs.
LEAD_ALIASES = {
    "Cybeam": "CyBeam", "Whitebox Branding": "Whitebox", "JustPlay": "Just Play Apps",
    "Dream Group": "Dream", "AEye Health": "AEYE Health", "Element": "Element Security",
    "Anchor 02": "Anchor", "Anchor 03": "Anchor", "SecuriThings": "SecuriThings",
    "Chamelio": "Chamelio", "Swimm": "Swimm", "Variant": "Variant", "Freehand": "Freehand",
    "Unibeam": "Unibeam",
}
# Several tabs are versions of one quote: tab -> (new-lead key, version order).
VERSION_OF = {
    "Air": ("Air", 1), "AIR1": ("Air", 2), "Copy of AIR1": ("Air", 3),
    "Anchor 02": ("Anchor", 1), "Anchor 03": ("Anchor", 2),
}
# New leads, traced from Nitsan's Gmail headers on 2026-10-03 (headers only).
# stage: 'won' links the client when one exists; 'lost' = outcome not recorded.
TRACED = {
    "Air": dict(company="Air", contact=("Yael Zarfati", "yael@air.security"), domain="air.security", first="2025-01-15", stage="won", client="Air"),
    "AIDOC": dict(company="Aidoc", contact=("Stav Adler", "stav.adler@aidoc.com"), domain="aidoc.com", first="2026-03-09", stage="won", client="AiDoc"),
    "CodeFlash": dict(company="CodeFlash", contact=("Saurabh Misra", "saurabh@codeflash.ai"), domain="codeflash.ai", first="2025-08-21", stage="won", client="CodeFlash"),
    "control monkey": dict(company="ControlMonkey", contact=("Tal Yizraeli", "tal@controlmonkey.io"), domain="controlmonkey.io", first="2025-10-13", stage="won", client="ControlMonkey"),
    "Intail": dict(company="Intail", contact=("Din Golan", "din@intail.ai"), domain="intail.ai", first="2025-05-27", stage="won", client=None),
    "ClearML": dict(company="ClearML", contact=("Noam Harel", "noamh@clearml.ai"), domain="clearml.ai", first="2026-01-14", stage="lost"),
    # Kept OPEN (Nitsan, 2026-10-03): mail as recent as July 2026; status updated in the app.
    "Getlira": dict(company="Getlira", contact=("Dan Gorfung", "dan@getlira.ai"), domain="getlira.ai", first="2026-02-23", last="2026-07-26", stage="open"),
    "Griiip": dict(company="Griiip", contact=("Tamir Plachinsky", "tamirp@griiip.com"), domain="griiip.com", first="2026-01-06", stage="lost"),
    "Jewish Climate Trust": dict(company="Jewish Climate Trust", contact=("Sarah Kandel-Finn", "sarah.kf@jct.org"), domain="jct.org", first="2025-10-23", stage="lost"),
    "Quack.ai": dict(company="Quack.ai", contact=("Shahar Cohen", "shahar@quack.ai"), domain="quack.ai", first="2025-05-08", stage="lost"),
    "Rokka.ai": dict(company="Rokka.ai", contact=("Nina Banai", "ninabanai@rokka.ai"), domain="rokka.ai", first="2025-03-02", stage="lost"),
    "Fattal": dict(company="Fattal", contact=("Maya Shalev", "mayash@fattal.co.il"), domain="fattal.co.il", first="2026-07-22", last="2026-07-22", stage="open"),
    "Itay Boneh": dict(company="Itay Boneh", contact=("Itay Boneh", "itaybo@gmail.com"), domain=None, first="2025-05-25", stage="lost"),
    # The tab says "Fizkal" — a typo for Fizikal (Nitsan). Came in through the website form.
    "Fizkal": dict(company="Fizikal", contact=("Sapir", "sapirb@movement-group.com"), domain="movement-group.com", first="2026-02-25", stage="lost"),
    "Dolphin": dict(company="Dolphin", contact=None, domain=None, first=None, stage="lost"),
}

NUM = (int, float)
IS_PERCENTABLE = re.compile(r"^(mobile|prep(are)?\s*(for|4)\s*dev\w*|qa)\b", re.I)
IS_DEV = re.compile(r"webflow|wordpress|development|\bdev\b|mcp", re.I)


def num(v):
    if isinstance(v, NUM) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        s = v.replace(",", "").strip()
        try:
            return float(s)
        except ValueError:
            return None
    return None


def text(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def category(name, phase):
    own = name.lower()
    ph = (phase or "").lower()
    if IS_PERCENTABLE.search(name):
        return "website"
    if is_dev(name):
        return "other"
    if re.search(r"strategy|workshop|alignment", own):
        return "strategy"
    if re.search(r"website|web design|\bpages?\b|app design", ph):
        return "website"
    if re.search(r"deck|slide|pitch|one.?pager|presentation", own):
        return "deck"
    if re.search(r"brand|visual|logo|identity|book|guideline", own):
        return "brand"
    if re.search(r"homepage|home page|page|404|legal|cookie|menu|footer|wireframe|blog|about|career|contact|demo|pricing|product|solution|resource|use case|platform|login|sign up|search|news|partner|customer|integration|security|feature", own):
        return "website"
    return "other"


def is_dev(name):
    """Development work — but not "Visual Language Development" or "Prep for Dev"."""
    if IS_PERCENTABLE.search(name) or re.search(r"visual|brand|logo", name, re.I):
        return False
    return bool(IS_DEV.search(name))


def pct_from(name, note_val):
    m = re.search(r"\((\d+(?:\.\d+)?)\s*%\)", name)
    if m:
        return float(m.group(1))
    if note_val is not None and 0 < note_val < 1:
        return round(note_val * 100, 2)
    return None


SKIP_NAME = re.compile(r"total|hourly rate|usd/ils|\(usd\)|^in (hrs|ils|usd)$|^ils|rate as of|^cap:", re.I)


def find_header(rows):
    for i, r in enumerate(rows[:6]):
        cells = [text(c).lower() for c in r]
        if any(re.search(r"\bmin\b|min hrs|min hours|^hours$|work hours", c) for c in cells):
            return i
    return None


def parse_rate(rows):
    for r in rows:
        name = text(r[0]) if r else ""
        joined = " ".join(text(c) for c in r)
        m = re.search(r"(\d{3})\s*ils", joined, re.I)
        if m:
            return float(m.group(1)), None
        if re.search(r"hourly rate", name, re.I):
            for c in r[1:]:
                v = num(c)
                if v:
                    return v, None
        m = re.search(r"(\d+)\s*usd per hr", joined, re.I)
        if m:
            return 350.0, f"originally quoted at {m.group(1)} USD/h — rate set to 350 NIS"
    return 350.0, "no rate stated in the tab — 350 assumed"


def sheet_total(rows, cmin, cmax):
    """The tab's own grand total, preferring the WHOLE-estimate rows over part totals."""
    prefs = [
        r"grand total in work hours",
        r"^total in work hours$",
        r"total with option 1 \(hrs\)",
        r"subtotal in hours",
        r"^total in work hours",
    ]
    for pat in prefs:
        for r in rows:
            name = text(r[0]) if r else ""
            if not re.search(pat, name, re.I):
                continue
            lo = num(r[cmin]) if cmin < len(r) else None
            hi = num(r[cmax]) if cmax is not None and cmax < len(r) else lo
            if lo is not None:
                return lo, hi if hi is not None else lo
    return None


def parse_tab(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    h = find_header(rows)
    if h is None:
        return None
    header = [text(c).lower() for c in rows[h]]
    # name column: the first header that names the item
    name_col = next((i for i, c in enumerate(header) if c in ("page", "item", "task", "deliverable", "column 1")), 0)
    if "deliverable" in header:
        name_col = header.index("deliverable")
    cmin = next(i for i, c in enumerate(header) if re.search(r"\bmin\b|min hrs|min hours|^hours$", c))
    cmax = next((i for i, c in enumerate(header) if re.search(r"\bmax\b|max hrs|max hours", c)), None)
    cnote = next((i for i, c in enumerate(header) if c in ("notes", "comments", "description")), None)
    cdesc = header.index("description") if "description" in header else None
    pkg_col = header.index("package") if "package" in header else None
    rate, rate_note = parse_rate(rows)
    total = sheet_total(rows[h + 1 :], cmin, cmax)

    body = rows[h + 1 :]
    items = []  # (kind, name, lo, hi, note_text, note_num, child, pkg)
    after_total = False
    for r in body:
        if text(r[0]).lower().startswith("total in work hours"):
            after_total = True
            continue
        if after_total and ws.title.strip() == "Intail":
            # The dev breakdown under the total sits one column to the left:
            # name | min | max. It is its own phase.
            nm, a, b = text(r[0]), num(r[1]) if len(r) > 1 else None, num(r[2]) if len(r) > 2 else None
            if nm and a is not None and not re.fullmatch(r"[\d.]+", nm):
                if not any(it[0] == "pkg" and it[1] == "Development" for it in items):
                    items.append(("pkg", "Development", None, None, "", None, False, "Development"))
                items.append(("row", nm, a, b if b is not None else a, "", None, False, "Development"))
            continue
        g = lambda i: r[i] if i is not None and i < len(r) else None
        raw = text(g(name_col))
        pkg = text(g(pkg_col)) if pkg_col is not None else ""
        lo, hi = num(g(cmin)), num(g(cmax)) if cmax is not None else None
        if hi is None and lo is not None:
            hi = lo
        note_v = g(cnote)
        note_num = num(note_v)
        note_txt = "" if note_num is not None else text(note_v)
        if cdesc is not None and cdesc != cnote:
            d = text(g(cdesc))
            note_txt = (d + (" — " + note_txt if note_txt else "")).strip()
        child = raw.startswith("↪")
        name = re.sub(r"^↪\s*", "", raw).strip()
        if not name and not pkg:
            continue
        header_total = re.search(r"\btotal$", name, re.I) and not re.search(r"work hours|in ils|in usd|hrs\)", name, re.I)
        if (SKIP_NAME.search(name) and not header_total) or (not name and pkg and lo is None):
            if pkg and not SKIP_NAME.search(pkg):
                items.append(("pkg", pkg, None, None, "", None, False, pkg))
            continue
        if re.fullmatch(r"\d+(\.0)?", name):  # "404.0"
            name = "404"
        if pkg and not pkg.startswith("*"):
            items.append(("pkg", pkg, None, None, "", None, False, pkg))
        items.append(("row", name, lo, hi, note_txt, note_num, child, pkg))

    # Build phases.
    #
    # ⚠️ TWO "CURRENT" PHASES. Children (↪ rows) go into the phase their
    # header opened; top-level rows with hours go into a "Scope" phase (or
    # "Development"). Keeping them apart is what stops a stray top-level line
    # in the middle of a website block (Air's "Custom Video section") from
    # swallowing the website pages that follow it.
    #
    # ⚠️ A TOP-LEVEL ROW FOLLOWED BY CHILDREN IS A HEADER ONLY IF ITS NUMBERS
    # ARE THE SUM OF THOSE CHILDREN (or it has none). Otherwise it is a real
    # line that happens to sit above a block — Air's "wireframes + sitemap".
    phases = []

    def new_phase(name, auto=False):
        ph = dict(name=name, lines=[], auto=auto)
        phases.append(ph)
        return ph

    child_phase = None
    top_phase = None

    def children_sum(j):
        lo = hi = 0.0
        k = j + 1
        while k < len(items) and items[k][0] == "row" and (items[k][6] or items[k][2] is None):
            if items[k][6] and items[k][2] is not None:
                lo += items[k][2]
                hi += items[k][3] or items[k][2]
            k += 1
        return lo, hi, k > j + 1

    def is_header(j):
        kind, name, lo, hi, *_ = items[j]
        nxt = items[j + 1] if j + 1 < len(items) else None
        if not nxt or nxt[0] != "row":
            return False
        # Only a DIRECT ↪ row makes a heading. Fattal's "careers" line sits above
        # "App design" (a heading of its own) — looking past it made "careers"
        # swallow App design's pages.
        has_child = nxt[6]
        if not has_child:
            return False
        if lo is None:
            return True
        clo, chi, _ = children_sum(j)
        n_children = 0
        k = j + 1
        while k < len(items) and items[k][0] == "row" and (items[k][6] or items[k][2] is None):
            n_children += 1 if items[k][6] else 0
            k += 1
        # A subtotal row: its hours are (roughly) its children's. Fizikal's
        # "Marketing Assets 104–160" over children summing 120–192 is still a
        # heading — the sheet's own subtotal was off — while Air's
        # "wireframes + sitemap 16–28" over ~400h of pages is plainly a line.
        return n_children >= 2 and clo > 0 and lo >= 0.5 * clo



    def make_line(name, lo, hi, note, note_num, phase_name):
        pct = pct_from(name, note_num) if IS_PERCENTABLE.search(name) else None
        line = dict(
            name=re.sub(r"\s*\(\d+%\)", "", name).strip(),
            description=note or None,
            min=lo,
            max=hi,
            category=category(name, phase_name),
            percent=pct,
            sheet=(lo, hi),
        )
        m = re.match(r"option\s*(\d+)\s*[–\-:]\s*(.+)", line["name"], re.I)
        if m:
            line["name"] = m.group(2).strip()
            line["alt_group"] = "Logo" if re.search(r"logo", line["name"], re.I) else phase_name
            line["chosen"] = m.group(1) == "1"
        return line

    i = 0
    while i < len(items):
        kind, name, lo, hi, note, note_num, child, pkg = items[i]
        nxt = items[i + 1] if i + 1 < len(items) else None
        if kind == "pkg":
            if not top_phase or top_phase["name"] != name:
                top_phase = new_phase(name)
                child_phase = top_phase
            i += 1
            continue
        if child:
            if lo is None:
                i += 1
                continue
            if child_phase is None:
                child_phase = new_phase("Website Design")
            child_phase["lines"].append(make_line(name, lo, hi, note, note_num, child_phase["name"]))
            i += 1
            continue
        # top level
        if re.search(r"\boption \d\b", name, re.I) and lo is not None and not (nxt and nxt[6]):
            i += 1  # "Brand Identity option 1" — a subtotal of one alternative path
            continue
        if is_header(i) or (lo is None and nxt and nxt[0] == "row" and nxt[6]):
            label = re.sub(r"\s+option\s*\d+$", "", re.sub(r"\s*total$", "", name, flags=re.I), flags=re.I).strip()
            child_phase = new_phase(label or "Website Design")
            i += 1
            continue
        if lo is None or re.search(r"\btotal\b", name, re.I):
            i += 1
            continue
        cat = category(name, None)
        if is_dev(name):
            target = next((p for p in phases if p["name"] == "Development"), None) or new_phase("Development", auto=True)
        elif cat == "website" and child_phase is not None and not child_phase.get("pkg"):
            target = child_phase  # a page or Mobile/QA row written at top level, under a website block
        elif top_phase is not None and (top_phase["auto"] or top_phase is child_phase):
            target = top_phase
        else:
            target = top_phase = new_phase("Scope", auto=True)
        if cat == "website" and target is top_phase and target["name"] == "Scope" and IS_PERCENTABLE.search(name) is None:
            pass
        target["lines"].append(make_line(name, lo, hi, note, note_num, target["name"]))
        i += 1

    # ⚠️ A PERCENTAGE ONLY STAYS A PERCENTAGE IF IT REPRODUCES THE SHEET'S OWN
    # HOURS. Getlira's "0.12" sits beside hours that are 15% of its pages, and
    # turning that into a live 12% would silently re-price the quote. Where the
    # two disagree, the sheet's hours win and the line stays fixed.
    est_lines = [l for p in phases for l in p["lines"]]
    # Iterate until stable: reverting one line to fixed hours changes the base
    # every other percentage is taken from (the app counts fixed-hour website
    # lines in the base — see lineHours in src/lib/leads/estimate.ts).
    changed = True
    while changed:
        changed = False
        base_lo = sum(l["min"] for l in est_lines if l["percent"] is None and l["category"] == "website")
        base_hi = sum(l["max"] for l in est_lines if l["percent"] is None and l["category"] == "website")
        for l in est_lines:
            if l["percent"] is None:
                continue
            lo_c, hi_c = base_lo * l["percent"] / 100, base_hi * l["percent"] / 100
            slo, shi = l["sheet"]
            if abs(lo_c - slo) > max(1.5, 0.06 * slo) or abs(hi_c - shi) > max(1.5, 0.06 * shi):
                l["percent"] = None  # keep the sheet's fixed hours
                changed = True

    phases = [p for p in phases if p["lines"]]
    return dict(rate=rate, rate_note=rate_note, phases=phases, sheet_total=total)


# Where the tab's total row is not the total of what is imported.
EXPECTED = {
    # Main table only: collateral 236–396 + website incl. dev 688.2–1376.4 + brand book 40–60.
    # The sheet's "Total in work Hours" also adds the side package (imported separately).
    "Freehand": (964.2, 1832.4),
    "Freehand [side]": (462.8, 761.8),
    # The breakdown's total excludes the development section (72–112) shown under it.
    "Intail": (288.0, 516.0),
}
# Tabs whose own total row does not add up their own lines — imported as the lines say.
SHEET_ARITHMETIC = {"Variant", "Fizkal"}  # Fizikal: "Marketing Assets" says 104–160, its items add up to 120–192


def totals(est):
    """Same rules as src/lib/leads/estimate.ts."""
    lines = [l for p in est["phases"] for l in p["lines"]]
    winners = {}
    for l in lines:
        g = l.get("alt_group")
        if g and l.get("chosen", True) and g not in winners:
            winners[g] = id(l)
    counts = lambda l: (not l.get("alt_group")) or winners.get(l["alt_group"]) == id(l)
    base_lo = sum(l["min"] for l in lines if l["percent"] is None and l["category"] == "website" and counts(l))
    base_hi = sum(l["max"] for l in lines if l["percent"] is None and l["category"] == "website" and counts(l))
    lo = hi = 0.0
    for l in lines:
        if not counts(l):
            continue
        if l["percent"] is not None:
            import math

            lo += math.ceil(base_lo * l["percent"] / 100 * 2) / 2
            hi += math.ceil(base_hi * l["percent"] / 100 * 2) / 2
        else:
            lo += l["min"]
            hi += l["max"]
    return lo, hi


# ── control monkey has Old / New side by side: two versions ───────────────────
def parse_control_monkey(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    out = []
    for lo_c, hi_c, label in ((2, 3, "Old"), (4, 5, "New")):
        lines = []
        for r in rows[2:]:
            name = text(r[1]) if len(r) > 1 else ""
            if not name or SKIP_NAME.search(name):
                continue
            lo, hi = num(r[lo_c]) if len(r) > lo_c else None, num(r[hi_c]) if len(r) > hi_c else None
            if lo is None:
                continue
            lines.append(dict(name=name, description=None, min=lo, max=hi if hi is not None else lo, category=category(name, None), percent=None))
        if lines:
            est = dict(rate=300.0, rate_note="ILS + VAT total implies 300/h", phases=[dict(name="Scope", lines=lines)], sheet_total=None, label=label)
            out.append(est)
    # sheet totals: Old 650/1090, New 508/652
    out[0]["sheet_total"] = (650.0, 1090.0)
    out[1]["sheet_total"] = (508.0, 652.0)
    return out


# ── Freehand's side table: a second, smaller package beside the main one ─────
def parse_freehand_side(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    start = next(i for i, r in enumerate(rows) if len(r) > 5 and text(r[5]).lower() == "item")
    lines = []
    for r in rows[start + 1 :]:
        name = text(r[5]) if len(r) > 5 else ""
        if not name:
            continue
        if re.match(r"in (hrs|ils|usd)", name, re.I):
            break
        lo, hi = num(r[6]) if len(r) > 6 else None, num(r[7]) if len(r) > 7 else None
        if lo is None:
            continue
        note = r[8] if len(r) > 8 else None
        nn = num(note)
        clean = re.sub(r"^↪\s*", "", name).strip()
        # Fixed hours, as the sheet has them: this package's 10% is of its
        # homepage + video only, a base the app's "% of website" can't express.
        pct = None
        lines.append(dict(name=clean, description=None if nn is not None else (text(note) or None), min=lo, max=hi if hi is not None else lo,
                          category=category(clean, None), percent=pct, sheet=(lo, hi)))
    return dict(rate=350.0, rate_note=None, phases=[dict(name="Brand identity & launch package", lines=lines, auto=False)],
                sheet_total=None, label="side")


# ── REST ───────────────────────────────────────────────────────────────────────
def rest(method, path, body=None, prefer=None):
    req = urllib.request.Request(f"{URL}/rest/v1/{path}", method=method)
    req.add_header("apikey", KEY)
    req.add_header("Authorization", f"Bearer {KEY}")
    req.add_header("Content-Type", "application/json")
    if prefer:
        req.add_header("Prefer", prefer)
    data = json.dumps(body).encode() if body is not None else None
    with urllib.request.urlopen(req, data) as r:
        raw = r.read()
        return json.loads(raw) if raw else None


def main():
    wb = openpyxl.load_workbook(PATH, data_only=True)
    plan = []  # (tab, est)
    for ws in wb.worksheets:
        if ws.title.strip() in SKIP_TABS:
            continue
        if ws.title.strip() == "control monkey":
            for est in parse_control_monkey(ws):
                plan.append((ws.title.strip(), est))
            continue
        est = parse_tab(ws)
        if est:
            plan.append((ws.title.strip(), est))
        if ws.title.strip() == "Freehand":
            plan.append(("Freehand", parse_freehand_side(ws)))

    if ONLY:
        plan = [(t, e) for t, e in plan if t == ONLY]
    print(f"\n{len(plan)} estimates from {len(wb.worksheets)} tabs (Monday skipped)\n")
    bad = 0
    for tab, est in plan:
        lo, hi = totals(est)
        key = tab + (f" [{est['label']}]" if est.get("label") in ("side",) else "")
        st = EXPECTED.get(key) or est["sheet_total"]
        n = sum(len(p["lines"]) for p in est["phases"])
        pct = sum(1 for p in est["phases"] for l in p["lines"] if l["percent"] is not None)
        alt = sum(1 for p in est["phases"] for l in p["lines"] if l.get("alt_group"))
        flag = ""
        if st:
            gap = max(abs(lo - st[0]), abs(hi - st[1]))
            tol = max(2, 0.01 * st[1])  # percentage lines round UP to the half hour
            if gap <= tol:
                flag = "  ✓ matches sheet" + (" (rounding)" if gap > 2 else "")
            elif tab in SHEET_ARITHMETIC:
                flag = f"  ≈ sheet's total row says {st[0]:g}–{st[1]:g}, but its own lines add up to this"
            else:
                flag = f"  ⚠ sheet says {st[0]:g}–{st[1]:g} (gap {gap:g}h)"
                bad += 1
        else:
            flag = "  (no total row to check)"
        label = f" [{est['label']}]" if est.get("label") else ""
        print(f"{tab+label:28} rate {est['rate']:g}  {len(est['phases'])} phases, {n} lines ({pct} %, {alt} options)  →  {lo:g}–{hi:g} h{flag}")
        if est.get("rate_note"):
            print(f"{'':28} note: {est['rate_note']}")
        for p in est["phases"]:
            print(f"{'':28}   · {p['name'][:40]}: " + ", ".join(
                (l['name'][:22] + (f" {l['percent']:g}%" if l['percent'] is not None else f" {l['min']:g}-{l['max']:g}") + (" [opt]" if l.get('alt_group') else ""))
                for l in p["lines"][:6]) + (" …" if len(p["lines"]) > 6 else ""))
    print(f"\n{bad} tab(s) with a total that does not match the sheet.\n")

    if not WRITE:
        print("DRY RUN — nothing written. Re-run with --write.\n")
        return
    write(plan)


def write(plan):
    stages = rest("GET", "lead_stages?select=id,kind,rule_key")
    won = next(s["id"] for s in stages if s["kind"] == "won")
    lost = next(s["id"] for s in stages if s["kind"] == "lost")
    offer_sent = next(s["id"] for s in stages if s.get("rule_key") == "offer_sent")
    clients = rest("GET", "clients?select=id,name")
    leads_cache = {}

    def lead_for(tab):
        key = VERSION_OF.get(tab, (tab,))[0]
        if key in leads_cache:
            return leads_cache[key]
        alias = LEAD_ALIASES.get(tab)
        if alias:
            found = rest("GET", f"leads?select=id,company,stage_id,created_at&company=ilike.{urllib.parse.quote(alias)}&order=created_at.desc")
            # Prefer a Won lead (Swimm has a won and a lost one), else the newest.
            pick = next((l for l in found if l["stage_id"] == won), found[0] if found else None)
            if pick:
                leads_cache[key] = (pick["id"], pick["company"])
                return leads_cache[key]
        t = TRACED.get(key) or TRACED.get(tab)
        ref = f"estimates:{key}"
        existing = rest("GET", f"leads?select=id,company&sheet_ref=eq.{urllib.parse.quote(ref)}")
        if existing:
            leads_cache[key] = (existing[0]["id"], existing[0]["company"])
            return leads_cache[key]
        client_id = None
        if t.get("client"):
            client_id = next((c["id"] for c in clients if c["name"].lower() == t["client"].lower()), None)
        when = f"{t['first']}T12:00:00+03:00" if t.get("first") else datetime.now(timezone.utc).isoformat()
        last = f"{t['last']}T12:00:00+03:00" if t.get("last") else when
        stage_id = {"won": won, "lost": lost, "open": offer_sent}[t["stage"]]
        row = dict(
            company=t["company"], domain=t.get("domain"), website=f"https://{t['domain']}" if t.get("domain") else None,
            stage_id=stage_id, sheet_ref=ref, client_id=client_id, source="website" if tab == "Fizkal" else None,
            won_at=when if t["stage"] == "won" else None,
            lost_note="Imported from the Estimations sheet — outcome not recorded" if t["stage"] == "lost" else None,
            created_at=when, stage_changed_at=last, last_activity_at=last,
        )
        made = rest("POST", "leads", row, prefer="return=representation")[0]
        if t.get("contact"):
            rest("POST", "lead_contacts", dict(lead_id=made["id"], name=t["contact"][0], email=t["contact"][1], position=1))
        rest("POST", "lead_events", [dict(lead_id=made["id"], kind="created", at=when,
            body=f"Imported from the Estimations / Pricing sheet (tab “{tab}”). Contact traced from Gmail headers.")])
        leads_cache[key] = (made["id"], made["company"])
        return leads_cache[key]

    import urllib.parse  # noqa: E402

    order = sorted(plan, key=lambda x: (VERSION_OF.get(x[0], (x[0], 1))[1], 0 if x[1].get("label") in (None, "Old") else 1))
    made = 0
    for tab, est in order:
        lead_id, company = lead_for(tab)
        note = f"Imported from sheet tab “{tab}”" + (f" ({est['label']})" if est.get("label") else "")
        # ⚠️ PREFIX match: a tab with a rate note stores "<note> · <rate note>",
        # which an exact match never finds — a re-run duplicated those.
        dup = rest("GET", f"lead_estimates?select=id&lead_id=eq.{lead_id}&change_note=like.{urllib.parse.quote(note + '*')}")
        if dup:
            continue
        last = rest("GET", f"lead_estimates?select=version&lead_id=eq.{lead_id}&order=version.desc&limit=1")
        version = (last[0]["version"] if last else 0) + 1
        e = rest("POST", "lead_estimates", dict(
            lead_id=lead_id, version=version, status="approved", rate=est["rate"], vat_percent=18,
            change_note=note + (f" · {est['rate_note']}" if est.get("rate_note") else ""),
            approved_at=datetime.now(timezone.utc).isoformat(),
        ), prefer="return=representation")[0]
        pos = 0
        for pi, p in enumerate(est["phases"], 1):
            ph = rest("POST", "estimate_phases", dict(estimate_id=e["id"], name=p["name"][:120], position=pi), prefer="return=representation")[0]
            for l in p["lines"]:
                pos += 1
                rest("POST", "estimate_lines", dict(
                    estimate_id=e["id"], phase_id=ph["id"], name=l["name"][:200], description=l["description"],
                    category=l["category"], kind="percent" if l["percent"] is not None else "hours",
                    min_hours=None if l["percent"] is not None else l["min"],
                    max_hours=None if l["percent"] is not None else l["max"],
                    percent=l["percent"], percent_of="website" if l["percent"] is not None else None,
                    alt_group=l.get("alt_group"), chosen=l.get("chosen", True), position=pos,
                ))
        made += 1
        print(f"  ✓ {company} v{version} ← {tab}")
    print(f"\nimported {made} estimates.\n")


if __name__ == "__main__":
    main()
